"""Host-side archive service for RetailPrime.

Run this on the Windows machine that owns Excel.  Browser and mobile clients
never access a Windows path: they call this HTTP service only after Firestore
has acknowledged a write.  The service upserts by stable document ID and never
recreates the master workbook.
"""
from __future__ import annotations

import json
import os
import threading
import urllib.request
import urllib.parse
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(os.getenv("AIOH_EXCEL_PATH", ROOT / "Excel" / "AIOH_ERP_Live.xlsx.xlsx"))
INVOICE_ROOT = Path(os.getenv("AIOH_INVOICE_ROOT", ROOT / "Invoice"))
ARCHIVE_KEY = os.getenv("AIOH_ARCHIVE_KEY", "")
LOCK = threading.Lock()
app = FastAPI(title="AIOH Local Archive")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173"], allow_methods=["POST", "GET"], allow_headers=["*"])

# Existing master-workbook sheets.  Reusing these preserves the archive layout;
# only genuinely new archive categories get a new sheet.
SHEETS = {"inventoryMovements": "InventoryArchive", "invoices": "SalesInvoices", "customers": "Customers", "products": "Products", "sales": "SalesInvoices", "payments": "Payments", "expenses": "Expenses", "services": "ServiceRecords", "coupons": "Coupons", "staff": "StaffAttendance", "attendance": "StaffAttendance", "settings": "Settings"}
LISTENERS: list[Any] = []
DATABASE_BACKEND = os.getenv("RETAILPRIME_DATABASE_BACKEND", "firestore")
if DATABASE_BACKEND not in {"firestore", "neon"}:
    raise RuntimeError("RETAILPRIME_DATABASE_BACKEND must be firestore or neon")
ARCHIVE_STATE: dict[str, Any] = {"backend": DATABASE_BACKEND, "status": "idle"}
STOP_ARCHIVE = threading.Event()

def require_key(key: str | None) -> None:
    if DATABASE_BACKEND == "neon" and not ARCHIVE_KEY:
        raise HTTPException(503, "Configure the server archive key before archiving")
    if ARCHIVE_KEY and key != ARCHIVE_KEY:
        raise HTTPException(401, "Invalid archive key")

def excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return json.dumps(value, ensure_ascii=False, default=str)

def upsert(sheet_name: str, record: dict[str, Any]) -> None:
    if not WORKBOOK.exists():
        raise FileNotFoundError(f"Master workbook not found: {WORKBOOK}")
    record_id = str(record.get("id") or record.get("invoiceId") or record.get("serviceId") or "").strip()
    if not record_id:
        raise ValueError("A stable record id is required")
    with LOCK:
        book = load_workbook(WORKBOOK)
        sheet = book[sheet_name] if sheet_name in book.sheetnames else book.create_sheet(sheet_name)
        headers = [cell.value for cell in sheet[1]] if sheet.max_row else []
        if not headers or not any(headers):
            headers = ["id", *[key for key in record if key != "id"]]
            for index, header in enumerate(headers, 1): sheet.cell(1, index, header)
        if "id" not in headers:
            headers.insert(0, "id")
            sheet.insert_cols(1)
            sheet.cell(1, 1, "id")
        for key in record:
            if key not in headers:
                headers.append(key); sheet.cell(1, len(headers), key)
        target = next((row for row in range(2, sheet.max_row + 1) if str(sheet.cell(row, 1).value or "") == record_id), sheet.max_row + 1)
        for index, header in enumerate(headers, 1):
            if header in record: sheet.cell(target, index, excel_value(record[header]))
        book.save(WORKBOOK)

def start_firestore_listener() -> None:
    """Load the preserved listener only for explicit Firestore mode."""
    from archive_firestore import start_firestore_listener as start_legacy
    start_legacy(ROOT, SHEETS, LISTENERS, upsert)


def neon_archive_query(connection: str, query: str, params: list[Any]) -> list[dict[str, Any]]:
    """Server-only credential, parameterized values and a fixed table allowlist."""
    url = urllib.parse.urlparse(connection)
    if not url.hostname or not url.hostname.endswith(".neon.tech") or url.path != "/neondb":
        raise ValueError("Invalid archive database configuration")
    request = urllib.request.Request(
        f"https://{url.hostname}/sql",
        data=json.dumps({"query": query, "params": params}).encode(),
        headers={"Content-Type": "application/json", "Neon-Connection-String": connection, "Neon-Array-Mode": "false"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=20) as response:
        return json.load(response).get("rows", [])


def start_neon_listener() -> None:
    # Supply a dedicated read-only archive connection, never a frontend secret.
    connection = os.getenv("AIOH_NEON_ARCHIVE_DATABASE_URL", "")
    owner = os.getenv("AIOH_ARCHIVE_BUSINESS_UID", "53RkfgSSN3USUvtceny57iiZiqj2")
    if not connection:
        ARCHIVE_STATE["status"] = "not_configured"
        return
    table_map = {"inventoryMovements": "inventory_movements", **{name: name for name in SHEETS if name != "inventoryMovements"}}
    created_only = {"inventoryMovements", "sales", "payments"}
    cursors: dict[str, tuple[str, str]] = {}

    def poll() -> None:
        while not STOP_ARCHIVE.is_set():
            try:
                for entity, table in table_map.items():
                    date_column = "created_at" if entity in created_only else "coalesce(updated_at,created_at)"
                    cursor_at, cursor_id = cursors.get(entity, ("-infinity", ""))
                    rows = neon_archive_query(connection, f"SELECT id,raw,coalesce({date_column},'-infinity'::timestamptz)::text AS cursor_at FROM retailprime.{table} WHERE owner_uid=$1 AND(coalesce({date_column},'-infinity'::timestamptz),id)>($2::timestamptz,$3) ORDER BY coalesce({date_column},'-infinity'::timestamptz),id LIMIT 250", [owner, cursor_at, cursor_id])
                    for row in rows:
                        upsert(SHEETS[entity], {**row["raw"], "id": row["id"]})
                        # Advance only after the workbook save succeeds.
                        cursors[entity] = (row["cursor_at"], row["id"])
                ARCHIVE_STATE["status"] = "running"
            except Exception:
                # Credentials, server responses and business payloads never
                # enter logs. A failed save retries; it never changes backend.
                ARCHIVE_STATE["status"] = "retry_pending"
            STOP_ARCHIVE.wait(30)

    worker = threading.Thread(target=poll, name="retailprime-neon-archive", daemon=True)
    LISTENERS.append(worker)
    worker.start()

@app.on_event("startup")
def startup() -> None:
    INVOICE_ROOT.joinpath("Bills").mkdir(parents=True, exist_ok=True)
    INVOICE_ROOT.joinpath("Services").mkdir(parents=True, exist_ok=True)
    INVOICE_ROOT.joinpath("Reports").mkdir(parents=True, exist_ok=True)
    if DATABASE_BACKEND == "neon":
        start_neon_listener()
    else:
        start_firestore_listener()


@app.on_event("shutdown")
def shutdown() -> None:
    STOP_ARCHIVE.set()
    for listener in LISTENERS:
        if hasattr(listener, "unsubscribe"):
            listener.unsubscribe()

@app.get("/health")
def health() -> dict[str, Any]:
    return {"workbook": str(WORKBOOK), "exists": WORKBOOK.exists(), "invoiceRoot": str(INVOICE_ROOT), "archive": ARCHIVE_STATE}

@app.post("/archive/{collection_name}")
def archive(collection_name: str, record: dict[str, Any], x_aioh_archive_key: str | None = Header(default=None)) -> dict[str, str]:
    require_key(x_aioh_archive_key)
    sheet = SHEETS.get(collection_name)
    if not sheet: raise HTTPException(404, "Unsupported archive collection")
    try:
        upsert(sheet, record)
    except (FileNotFoundError, ValueError) as error:
        raise HTTPException(422, str(error)) from error
    return {"status": "upserted", "sheet": sheet, "id": str(record["id"])}
